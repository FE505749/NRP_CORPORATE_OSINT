import pytest
from openpyxl import Workbook, load_workbook

from nrp_osint import denue_local, lote
from nrp_osint.config import Settings
from nrp_osint.persona import buscar_persona
from tests.test_api import cliente, esperar
from tests.test_denue_local import SUCURSALES, make_denue
from tests.test_dominio_y_formato import GRUPO


def excel(ruta, filas, titulos=lote.COLUMNAS):
    wb = Workbook()
    wb.active.append(titulos)
    for f in filas:
        wb.active.append(f)
    wb.save(ruta)
    return ruta


def ajustes(tmp_path):
    base = make_denue(tmp_path, SUCURSALES + GRUPO)
    denue_local.build(base)
    return Settings(db_path=str(tmp_path / "x.sqlite"), denue_local_path=str(denue_local.db_path(base)),
                    tukan_local_path="", enable_ddgs=False)


def test_leer_formato_y_validar_columnas(tmp_path):
    ruta = excel(tmp_path / "e.xlsx", [["1", "Ana", "", "Diaz", "Soto", "Profu Afore"], [None] * 6],
                 ["Cuenta", "Primer Nombre", "Segundo Nombre", "Primer Apellido", "Segundo Apellido", "EMPRESA"])
    assert lote.leer(ruta) == [{"cuenta": "1", "primer_nombre": "Ana", "segundo_nombre": "",
                                "primer_apellido": "Diaz", "segundo_apellido": "Soto", "empresa": "Profu Afore"}]
    with pytest.raises(ValueError, match="primer_apellido"):
        lote.leer(excel(tmp_path / "mal.xlsx", [["1", "Ana", "Profu"]], ["cuenta", "primer_nombre", "empresa"]))


def test_cada_fila_igual_que_la_busqueda_individual(tmp_path):
    s = ajustes(tmp_path)
    filas = [
        {"cuenta": "A1", "primer_nombre": "Francisco", "segundo_nombre": "Javier", "primer_apellido": "Espinoza",
         "segundo_apellido": "Mendieta", "empresa": "Profu Afore"},
        {"cuenta": "A2", "primer_nombre": "Francisco", "segundo_nombre": "Javier", "primer_apellido": "Espinoza",
         "segundo_apellido": "Mendieta", "empresa": "BancoX"},
        {"cuenta": "A3", "primer_nombre": "Ana", "segundo_nombre": "", "primer_apellido": "", "segundo_apellido": "",
         "empresa": "Profu Afore"},  # incompleta: sale en el Excel con su motivo
    ]
    res = lote.procesar(s, filas)
    assert [r["correo"] for r in res] == ["francisco.espinoza@profu.com.mx",
                                          "franciscojavier.espinoza@bancox.com.mx", ""]
    assert [r["patron"] for r in res] == ["first.paternal", "given.paternal", ""]
    assert res[2]["estado"] == "FALTAN_DATOS"
    individual = buscar_persona(s, "Francisco", "Javier", "Espinoza", "Mendieta", "BancoX")
    assert {k: individual[k] for k in ("correo", "dominio", "patron", "veredicto")} == \
           {k: res[1][k] for k in ("correo", "dominio", "patron", "veredicto")}

    destino = lote.exportar(res, tmp_path / "salida.xlsx")
    wb = load_workbook(destino)
    assert wb.sheetnames == ["Correos", "Detalle"]
    correos = list(wb["Correos"].iter_rows(values_only=True))
    assert correos[0] == ("Cuenta", "Correo estimado", "Patrón", "Dominio")
    assert correos[2] == ("A2", "franciscojavier.espinoza@bancox.com.mx", "nombresjuntos.apellido", "bancox.com.mx")
    detalle = list(wb["Detalle"].iter_rows(values_only=True))
    assert "Veredicto" in detalle[0] and len(detalle) == 4


def test_api_lote_y_plantilla(tmp_path):
    c = cliente(tmp_path)
    plantilla = c.get("/api/lote/plantilla")
    assert plantilla.status_code == 200
    ruta = excel(tmp_path / "subir.xlsx", [["Z1", "Ana", "", "Diaz", "Soto", "Profu Afore"]])
    r = c.post("/api/lote", files={"archivo": ("subir.xlsx", ruta.read_bytes(),
                                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    t = esperar(c, r.json()["trabajo"])
    assert t["estado"] == "listo", t["error"]
    assert t["resultado"]["con_correo"] == 1
    assert t["resultado"]["vista"][0] == {"cuenta": "Z1", "correo": "ana.diaz@profu.com.mx",
                                          "patron": "nombre.apellido", "dominio": "profu.com.mx"}
    assert c.get(f"/api/descargas/{t['resultado']['archivo']}").status_code == 200
    # Columnas equivocadas: error claro antes de empezar.
    mal = excel(tmp_path / "mal.xlsx", [["1", "Ana"]], ["cuenta", "nombre"])
    r = c.post("/api/lote", files={"archivo": ("mal.xlsx", mal.read_bytes(), "application/octet-stream")})
    assert r.status_code == 400 and "faltan columnas" in r.json()["detail"]
