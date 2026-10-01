from pathlib import Path
import csv,re
from openpyxl import load_workbook
from .normalize import clean_text,normalize_nrp,normalize_company,canonical_person,employer_key,person_key
from .tabular import csv_format,csv_incrustado

# Tuplas: el orden fija la prioridad de forma reproducible.
ALIASES={
 'account_id':('IDCUENTA','ID_CUENTA','CUENTA','ID CUENTA','ACCOUNT_ID','NO_CUENTA','NUM_CUENTA'),
 'nss':('NSS','NUMERO_SEGURIDAD_SOCIAL','NUM_SEG_SOCIAL'),
 'entity':('ENTIDAD','ESTADO','STATE'),
 'curp2':('CURP_2','CURP2','CURP'),
 'age':('EDAD','AGE'),
 'age_origin':('EDAD_ORIGEN','EDAD ORIGEN'),
 'nrp':('NRP','REGISTRO_PATRONAL','REG_PATRONAL','REGISTRO PATRONAL'),
 'company':('NOMBRE_EMPRESA','RAZON_SOCIAL','RAZÓN_SOCIAL','NOMBRE_ORIGINAL','NOMBRE_ORIGINAL_PADRON',
            'NOMBRE_COMERCIAL','NOMBRE_COMERCIAL_PADRON','EMPRESA','COMPANY','EMPLEADOR','PATRON','PATRÓN'),
 'blindaje':('BLINDAJE',),
 'full_name':('NOMBRE_COMPLETO','NOMBRE COMPLETO','FULL_NAME'),
 'first_names':('NOMBRE','NOMBRES','FIRST_NAME','FIRST_NAMES'),
 'paternal':('A_PATERNO','APELLIDO_PATERNO','PATERNO','LAST_NAME','APELLIDO'),
 'maternal':('A_MATERNO','APELLIDO_MATERNO','MATERNO','SECOND_LAST_NAME'),
 'sex':('SEXO','SEX'),
 'institute':('INSTITUTO','INSTITUTE'),
}
EXTERNAL_BLOCKED={'nss','curp2'}

def nh(x):return re.sub(r'[^A-Z0-9]+','_',str(x).strip().upper()).strip('_')
def map_columns(headers):
    normalized={nh(h):h for h in headers}; out={}; usadas=set()
    for field,names in ALIASES.items():
        for n in names:
            clave=nh(n)
            if clave in normalized and clave not in usadas:
                out[field]=normalized[clave];usadas.add(clave);break
    if 'account_id' not in out:
        raise ValueError('Falta columna requerida: account_id (cuenta / id_cuenta / IDCUENTA)')
    if 'first_names' not in out and 'full_name' not in out:
        raise ValueError('Falta columna requerida: nombre (nombre / nombres / NOMBRE_COMPLETO)')
    return out

def row_to_person(row,m,row_id,source):
    g=lambda f:clean_text(row.get(m.get(f,''),''))
    first,pat,mat=g('first_names'),g('paternal'),g('maternal')
    completo=g('full_name')
    if completo and not (first or pat or mat):
        partes=completo.split()
        first=' '.join(partes[:-2]) or (partes[0] if partes else '')
        pat=partes[-2] if len(partes)>=3 else (partes[-1] if len(partes)==2 else '')
        mat=partes[-1] if len(partes)>=2 else ''
        full=completo
    else:
        full=canonical_person(first,pat,mat)
    nrp=normalize_nrp(g('nrp')); company=g('company')
    return {'row_id':row_id,'account_id':g('account_id'),'person_key':person_key(full,g('entity')),'full_name':full,'first_names':first,'paternal':pat,'maternal':mat,
      'entity':g('entity'),'nrp':nrp,'company_input':company,'company_norm':normalize_company(company),'blindaje':g('blindaje'),'sex':g('sex'),'institute':g('institute'),'age':g('age'),'age_origin':g('age_origin'),
      'nss_local':g('nss'),'curp2_local':g('curp2'),'employer_key':employer_key(nrp,company),'ingest_source':source}

def _iter_csv(path):
    encoding,dialect=csv_format(path)
    with open(path,'r',encoding=encoding,newline='') as f:
        rd=csv.DictReader(f,dialect=dialect); m=map_columns(rd.fieldnames or [])
        for i,row in enumerate(rd,1):yield row_to_person(row,m,i,Path(path).name),m

def _iter_xlsx(path):
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        ws=wb.active; it=ws.iter_rows(values_only=True); headers=[clean_text(x) for x in next(it)]
        empacado=csv_incrustado(headers)
        if empacado:
            headers,dialecto=empacado; m=map_columns(headers); n=len(headers)
            for i,vals in enumerate(it,1):
                cruda=clean_text(vals[0] if vals else '')
                if not cruda:continue
                campos=next(csv.reader([cruda],dialecto))
                row={h:(campos[j] if j<n and j<len(campos) else '') for j,h in enumerate(headers)}
                yield row_to_person(row,m,i,Path(path).name),m
            return
        m=map_columns(headers); idx={h:i for i,h in enumerate(headers)}
        for i,vals in enumerate(it,1):
            row={h:(vals[j] if j<len(vals) else '') for h,j in idx.items()}; yield row_to_person(row,m,i,Path(path).name),m
    finally:
        wb.close()

def ingest(db,path,batch_size=5000):
    p=Path(path); iterator=_iter_csv(p) if p.suffix.lower()=='.csv' else _iter_xlsx(p); batch=[]; count=0; mapping={}
    for person,mapping in iterator:
        if not person['full_name']:continue
        batch.append(person);count+=1
        if len(batch)>=batch_size:db.upsert_people(batch);batch=[]
    if batch:db.upsert_people(batch)
    db.ensure_employer_jobs()
    return {'rows':count,'mapping':mapping,'external_search_excludes':['NSS','CURP_2']}
