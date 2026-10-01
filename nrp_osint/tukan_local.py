"""Copia local de Tukan: contactos por registro patronal (NRP).

Segunda fuente local despues de INEGI DENUE. El archivo trae NRP, EMPRESA,
ESTATUS y EMAIL; se guarda en data/tukan/tukan.sqlite con indice por NRP y
FTS5 sobre el nombre, y se consulta sin Internet.

Que aporta:
  - Empresa por NRP exacto, cuando la base del trabajador no trae el nombre.
  - Dominio corporativo, tomado del correo de contacto (nunca de correos
    gratuitos ni de dominios compartidos por despachos contables).
  - Patron de correo, a partir de los correos con forma de persona del dominio.
  - Correo de contacto del patron, prefiriendo ESTATUS = ACTIVO.
"""
from __future__ import annotations

import csv
import sqlite3
import time
from collections import Counter
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook
from rapidfuzz import fuzz

from .denue_local import FREE_MAIL, _fts_query, _score, denue_norm
from .domain import domain_company_score, mail_domain, root_domain
from .emailintel import pattern_from_local_part, weak_pattern_confidence
from .io_stream import nh
from .normalize import clean_text, normalize_nrp
from .paths import PROJECT_ROOT
from .tabular import csv_format

DEFAULT_DIR = PROJECT_ROOT / "data" / "tukan"
# Donde se busca el archivo si no se indica otro.
SOURCE_CANDIDATES = (DEFAULT_DIR / "Tukan.xlsx", PROJECT_ROOT / "Tukan.xlsx", PROJECT_ROOT.parent / "Tukan.xlsx")

ALIASES = {
    "nrp": ("NRP", "REGISTRO_PATRONAL", "REG_PATRONAL"),
    "empresa": ("EMPRESA", "RAZON_SOCIAL", "NOMBRE_EMPRESA", "PATRON"),
    "estatus": ("ESTATUS", "STATUS", "ESTADO"),
    "email": ("EMAIL", "CORREO", "E_MAIL", "CORREO_ELECTRONICO"),
}

# Proveedores de Internet y variantes que en Tukan aparecen como si fueran empresa.
ISP_MAIL = {
    "telmexmail.com", "megared.net.mx", "att.net.mx", "att.net", "prodigy.net", "eninfinitum.com", "ymail.com",
    "mail.com", "gmail.com.mx", "hotmail.com.mx", "outlook.com.mx", "rocketmail.com", "cablevision.net.mx",
    "todito.com", "latinmail.com", "starmedia.com", "usa.net", "sbcglobal.net", "comcast.net", "googlemail.com",
}
_FREE_STEMS = ("gmail", "hotmail", "yahoo", "outlook", "live", "prodigy", "icloud", "infinitummail")
_FREE_SUFFIXES = {"com", "com.mx", "es", "net", "net.mx", "mx"}

# Un dominio con correos de mas empresas que esto es de un tercero (despacho
# contable, outsourcing) salvo que el nombre de la empresa se parezca al dominio.
SHARED_LIMIT = 3
STATUS_RANK = {"ACTIVO": 0}

SCHEMA = """
PRAGMA journal_mode=OFF;
PRAGMA synchronous=OFF;
CREATE TABLE contacto(
  id INTEGER PRIMARY KEY,
  nrp TEXT, empresa TEXT, empresa_norm TEXT, estatus TEXT, email TEXT,
  dominio TEXT            -- dominio del correo si es corporativo, si no ''
);
CREATE TABLE dominio(dominio TEXT PRIMARY KEY, empresas INTEGER, nrps INTEGER, correos INTEGER);
CREATE TABLE meta(clave TEXT PRIMARY KEY, valor TEXT);
"""
POST = """
CREATE INDEX idx_contacto_nrp ON contacto(nrp);
CREATE INDEX idx_contacto_emp ON contacto(empresa_norm);
CREATE INDEX idx_contacto_dom ON contacto(dominio);
CREATE VIRTUAL TABLE fts USING fts5(empresa_norm, content='contacto', content_rowid='id', tokenize='unicode61');
INSERT INTO fts(fts) VALUES('rebuild');
INSERT INTO dominio
  SELECT dominio, COUNT(DISTINCT CASE WHEN empresa_norm<>'' THEN empresa_norm ELSE nrp END),
         COUNT(DISTINCT nrp), COUNT(DISTINCT email)
  FROM contacto WHERE dominio<>'' GROUP BY dominio;
ANALYZE;
"""


def db_path(base_dir: Path | str = DEFAULT_DIR) -> Path:
    return Path(base_dir) / "tukan.sqlite"


def default_source_file() -> Path | None:
    return next((p for p in SOURCE_CANDIDATES if p.exists()), None)


def is_free_mail(domain: str) -> bool:
    """Gratuitos, proveedores de Internet y errores de captura ('hotmai.com', 'gamil.com')."""
    if domain in FREE_MAIL or domain in ISP_MAIL:
        return True
    stem, _, suffix = domain.partition(".")
    return suffix in _FREE_SUFFIXES and len(stem) >= 4 and any(fuzz.ratio(stem, s) >= 80 for s in _FREE_STEMS)


def corporate_domain(email: str) -> str:
    if "@" not in email:
        return ""
    d = root_domain(email.split("@", 1)[1])
    return "" if not d or "." not in d or is_free_mail(d) else d


# --------------------------------------------------------------------------- carga

def _map_columns(headers) -> dict[str, int]:
    keys = {nh(h): i for i, h in enumerate(headers) if h}
    out = {}
    for field, names in ALIASES.items():
        for n in names:
            if nh(n) in keys:
                out[field] = keys[nh(n)]
                break
    faltan = [f for f in ("nrp", "email") if f not in out]
    if faltan:
        raise ValueError(f"Al archivo de Tukan le faltan columnas: {', '.join(faltan).upper()}")
    return out


def _iter_rows(path: Path):
    if path.suffix.lower() == ".csv":
        encoding, dialect = csv_format(path)
        with open(path, encoding=encoding, newline="") as f:
            it = csv.reader(f, dialect)
            m = _map_columns(next(it, []))
            for vals in it:
                yield m, vals
        return
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        it = wb.active.iter_rows(values_only=True)
        m = _map_columns([clean_text(x) for x in next(it, ())])
        for vals in it:
            yield m, vals
    finally:
        wb.close()


def build(source: Path | str, base_dir: Path | str = DEFAULT_DIR, progress=None) -> dict:
    """Construye tukan.sqlite desde el Excel o CSV. Escribe en un archivo temporal
    y lo reemplaza al final, asi una carga interrumpida no deja la base a medias."""
    src = Path(source)
    final = db_path(base_dir)
    final.parent.mkdir(parents=True, exist_ok=True)
    tmp = final.with_suffix(".building")
    tmp.unlink(missing_ok=True)
    conn = sqlite3.connect(tmp)
    conn.executescript(SCHEMA)
    total = activos = con_dominio = 0
    lote = []
    sql = "INSERT INTO contacto(nrp,empresa,empresa_norm,estatus,email,dominio) VALUES(?,?,?,?,?,?)"
    for m, vals in _iter_rows(src):
        g = lambda f: clean_text(vals[m[f]]) if f in m and m[f] < len(vals) else ""
        nrp, email = normalize_nrp(g("nrp")), g("email").lower()
        if not nrp or "@" not in email:
            continue
        empresa, estatus = g("empresa"), g("estatus").upper()
        dom = corporate_domain(email)
        activos += estatus == "ACTIVO"
        con_dominio += bool(dom)
        lote.append((nrp, empresa, denue_norm(empresa) if empresa else "", estatus, email, dom))
        if len(lote) >= 50000:
            conn.executemany(sql, lote)
            total += len(lote)
            lote = []
            if progress:
                progress(total)
    if lote:
        conn.executemany(sql, lote)
        total += len(lote)
    conn.commit()
    if progress:
        progress(total)
    conn.executescript(POST)
    nrps = conn.execute("SELECT COUNT(DISTINCT nrp) FROM contacto").fetchone()[0]
    dominios = conn.execute("SELECT COUNT(*) FROM dominio").fetchone()[0]
    compartidos = conn.execute("SELECT COUNT(*) FROM dominio WHERE empresas>?", (SHARED_LIMIT,)).fetchone()[0]
    meta = {"contactos": total, "nrp": nrps, "activos": activos, "con_dominio_empresa": con_dominio,
            "dominios": dominios, "dominios_compartidos": compartidos, "archivo": src.name,
            "construido": time.strftime("%Y-%m-%d %H:%M")}
    conn.executemany("INSERT INTO meta VALUES(?,?)", [(k, str(v)) for k, v in meta.items()])
    conn.commit()
    conn.execute("VACUUM")
    conn.close()
    final.unlink(missing_ok=True)
    tmp.replace(final)
    return meta


def info(base_dir: Path | str = DEFAULT_DIR) -> dict | None:
    p = db_path(base_dir)
    if not p.exists():
        return None
    try:
        conn = sqlite3.connect(f"file:{p.as_posix()}?mode=ro", uri=True)
        try:
            return {k: v for k, v in conn.execute("SELECT clave,valor FROM meta")}
        finally:
            conn.close()
    except sqlite3.Error:
        return None


# --------------------------------------------------------------------------- consulta

def _contact_rank(r) -> tuple:
    """Mejor contacto: ACTIVO antes que CADUCO/REVOCADO, corporativo antes que gratuito."""
    return (STATUS_RANK.get(r["estatus"], 1), not r["dominio"], r["id"])


class TukanLocalSource:
    name = "TUKAN_LOCAL"

    def __init__(self, path: Path | str | None = None, candidates: int = 300):
        self.path = Path(path) if path else db_path()
        self.conn = sqlite3.connect(f"file:{self.path.as_posix()}?mode=ro", uri=True, check_same_thread=False)
        self.conn.row_factory = sqlite3.Row
        self.candidates = candidates
        self._domains: dict[str, set] = {}

    @staticmethod
    def available(path: Path | str | None = None) -> bool:
        return (Path(path) if path else db_path()).exists()

    def _domain_stats(self, domain):
        r = self.conn.execute("SELECT empresas,nrps FROM dominio WHERE dominio=?", (domain,)).fetchone()
        return (r["empresas"], r["nrps"]) if r else (0, 0)

    def _summary(self, rows, company: str, match: str, score: float) -> dict:
        rows = sorted(rows, key=_contact_rank)
        best = rows[0]
        empresa = next((r["empresa"] for r in rows if r["empresa"]), "")
        name = company or empresa
        domain, domain_conf, shared = "", 0.0, ""
        votes = Counter(r["dominio"] for r in rows if r["dominio"])
        activos = {r["dominio"] for r in rows if r["dominio"] and r["estatus"] == "ACTIVO"}
        # Primero los dominios con algun contacto ACTIVO; a igualdad, el mas repetido.
        for cand in sorted(votes, key=lambda d: (d in activos, votes[d]), reverse=True):
            empresas, _ = self._domain_stats(cand)
            if empresas > SHARED_LIMIT and domain_company_score(cand, name) < .55:
                shared = shared or cand  # de un despacho: sirve como contacto, no como dominio
                continue
            activo = cand in activos
            domain = cand
            domain_conf = (.86 if activo else .80) * (1.0 if match == "NRP" else min(1.0, score))
            break
        return {
            "company": empresa, "nrp": best["nrp"], "match": match, "score": score,
            "domain": domain, "domain_confidence": round(domain_conf, 3), "shared_domain": shared,
            "email": best["email"], "status": best["estatus"],
            "emails": [(r["email"], r["estatus"]) for r in rows],
            "url": f"tukan-local://{best['nrp']}", "source": self.name,
        }

    def by_nrp(self, nrp: str, company: str = "") -> dict | None:
        nrp = normalize_nrp(nrp)
        rows = self.conn.execute("SELECT * FROM contacto WHERE nrp=?", (nrp,)).fetchall() if nrp else []
        if not rows:
            return None
        empresa = next((r["empresa_norm"] for r in rows if r["empresa_norm"]), "")
        target = denue_norm(company) if company else ""
        score = fuzz.WRatio(target, empresa) / 100 if target and empresa else 1.0
        return self._summary(rows, company, "NRP", score)

    def by_company(self, company: str) -> dict | None:
        """Busqueda por nombre (igualdad normalizada y luego FTS), como en DENUE local."""
        target = denue_norm(company)
        if not target:
            return None
        rows = self.conn.execute("SELECT * FROM contacto WHERE empresa_norm=? LIMIT ?",
                                 (target, self.candidates)).fetchall()
        if rows:
            return self._summary(rows, company, "NOMBRE", 1.0)
        for any_token in (False, True):
            q = _fts_query(target, any_token)
            if not q:
                return None
            try:
                rows = self.conn.execute(
                    "SELECT contacto.* FROM fts JOIN contacto ON contacto.id=fts.rowid WHERE fts MATCH ? "
                    "ORDER BY bm25(fts) LIMIT ?", (q, self.candidates)).fetchall()
            except sqlite3.OperationalError:
                rows = []
            if rows:
                break
        if not rows:
            return None
        scored = max(rows, key=lambda r: _score(target, r["empresa_norm"], ""))
        score = _score(target, scored["empresa_norm"], "")
        same = [r for r in rows if r["empresa_norm"] == scored["empresa_norm"]]
        return self._summary(same, company, "NOMBRE", score)

    def emails_for_domain(self, domain: str) -> set[str]:
        """Correos cuyo dominio es exactamente ese: lo que va despues de la @."""
        dom = mail_domain(domain)
        if not dom:
            return set()
        if dom not in self._domains:
            self._domains[dom] = {r[0] for r in self.conn.execute(
                "SELECT DISTINCT email FROM contacto WHERE dominio=?", (root_domain(dom),)) if mail_domain(r[0]) == dom}
        return self._domains[dom]

    def statuses_for_domain(self, domain: str) -> dict[str, str]:
        """{correo: ESTATUS} del dominio exacto. Si un correo aparece en varios
        patrones con estatus distinto, gana ACTIVO."""
        dom = mail_domain(domain)
        out: dict[str, str] = {}
        if not dom:
            return out
        for email, estatus in self.conn.execute(
                "SELECT email, estatus FROM contacto WHERE dominio=?", (root_domain(dom),)):
            if mail_domain(email) != dom:
                continue
            if email not in out or STATUS_RANK.get(estatus, 1) < STATUS_RANK.get(out[email], 1):
                out[email] = estatus
        return out

    def email_pattern_votes(self, domain: str) -> tuple[dict, int]:
        """Patrones de los correos con forma de persona del dominio. Sin nombre:
        se aceptan formas de una palabra solo si el resto es un apellido comun."""
        votes = Counter()
        for email in self.emails_for_domain(domain):
            pat = pattern_from_local_part(email.split("@", 1)[0], allow_single=True)
            if pat:
                votes[pat] += 1
        return dict(votes), sum(votes.values())

    def pattern(self, domain: str) -> tuple[dict, float]:
        votes, _ = self.email_pattern_votes(domain)
        return (votes, weak_pattern_confidence(votes, cap=.70)) if votes else ({}, 0.)


@lru_cache(maxsize=1)
def default_source() -> TukanLocalSource | None:
    return TukanLocalSource() if TukanLocalSource.available() else None
