"""Excel de personas: se lee, se busca cada fila igual que en la busqueda
individual (persona.buscar_personas) y se exporta un Excel con dos pestanas:

  Correos  cuenta, correo estimado, patron y dominio. Nada mas.
  Detalle  todo lo que explica cada correo: empresa encontrada, confianza,
           correos reales del dominio, desglose de formatos y veredicto.

Formato de entrada (primera fila = titulos, en la hoja activa):
  cuenta, primer_nombre, segundo_nombre, primer_apellido, segundo_apellido, empresa
"""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .io_stream import nh
from .normalize import clean_text
from .persona import ETIQUETA_ESTADO, ETIQUETA_FORMATO, ETIQUETA_RUTA, ETIQUETA_TIPO, buscar_personas
from .tabular import csv_format

COLUMNAS = ["cuenta", "primer_nombre", "segundo_nombre", "primer_apellido", "segundo_apellido", "empresa"]
OBLIGATORIAS = ["cuenta", "primer_nombre", "primer_apellido", "empresa"]
# Variantes que tambien se aceptan en el titulo (sin importar mayusculas ni espacios).
ALIAS = {
    "cuenta": ("CUENTA", "ID_CUENTA", "IDCUENTA", "NO_CUENTA", "NUM_CUENTA"),
    "primer_nombre": ("PRIMER_NOMBRE", "NOMBRE_1", "NOMBRE1"),
    "segundo_nombre": ("SEGUNDO_NOMBRE", "NOMBRE_2", "NOMBRE2"),
    "primer_apellido": ("PRIMER_APELLIDO", "APELLIDO_PATERNO", "A_PATERNO", "PATERNO"),
    "segundo_apellido": ("SEGUNDO_APELLIDO", "APELLIDO_MATERNO", "A_MATERNO", "MATERNO"),
    "empresa": ("EMPRESA", "NOMBRE_EMPRESA", "RAZON_SOCIAL"),
}

AZUL, BLANCO = "073056", "FFFFFF"


def _mapa(titulos: list[str]) -> dict[str, int]:
    pos = {nh(t): i for i, t in enumerate(titulos) if clean_text(t)}
    mapa = {}
    for campo, nombres in ALIAS.items():
        for n in nombres:
            if nh(n) in pos:
                mapa[campo] = pos[nh(n)]
                break
    faltan = [c for c in OBLIGATORIAS if c not in mapa]
    if faltan:
        raise ValueError("Al Excel le faltan columnas: " + ", ".join(faltan)
                         + ". El formato es: " + ", ".join(COLUMNAS) + ".")
    return mapa


def leer(ruta: Path | str) -> list[dict]:
    ruta = Path(ruta)
    if ruta.suffix.lower() == ".csv":
        encoding, dialecto = csv_format(ruta)
        with open(ruta, encoding=encoding, newline="") as f:
            filas = list(csv.reader(f, dialecto))
    else:
        wb = load_workbook(ruta, read_only=True, data_only=True)
        try:
            filas = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    if not filas:
        raise ValueError("El archivo está vacío.")
    mapa = _mapa([clean_text(t) for t in filas[0]])
    salida = []
    for vals in filas[1:]:
        fila = {c: clean_text(vals[i]) if i < len(vals) else "" for c, i in mapa.items()}
        if any(fila.values()):
            salida.append({c: fila.get(c, "") for c in COLUMNAS})
    if not salida:
        raise ValueError("El Excel no tiene filas con datos debajo de los títulos.")
    return salida


def procesar(settings, filas: list[dict], avance=None) -> list[dict]:
    return buscar_personas(settings, filas, avance=avance)


# --------------------------------------------------------------------------- Excel

def _hoja(ws, titulos: list[str], filas: list[list], anchos: list[int]):
    ws.append(titulos)
    for c in ws[1]:
        c.font = Font(bold=True, color=BLANCO)
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(vertical="center")
    for f in filas:
        # Un texto que empieza con '=' Excel lo tomaria como formula.
        ws.append(["'" + v if isinstance(v, str) and v[:1] in "=+@" else v for v in f])
    for i, ancho in enumerate(anchos):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = ancho
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def exportar(resultados: list[dict], destino: Path | str) -> Path:
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    patron = lambda r: ETIQUETA_FORMATO.get(r["patron"], r["patron"]) if r["patron"] else ""

    _hoja(wb.active, ["Cuenta", "Correo estimado", "Patrón", "Dominio"],
          [[r["cuenta"], r["correo"], patron(r), r["dominio"]] for r in resultados], [16, 44, 26, 30])
    wb.active.title = "Correos"

    detalle = []
    for r in resultados:
        reales = r["correos_reales"]
        detalle.append([
            r["cuenta"], r["primer_nombre"], r["segundo_nombre"], r["primer_apellido"], r["segundo_apellido"],
            r["empresa_entrada"], r["empresa"], r["correo"], patron(r), r["dominio"],
            ETIQUETA_TIPO.get(r["tipo"], r["tipo"]) if r["correo"] else ETIQUETA_ESTADO.get(r["estado"], r["estado"]),
            round(float(r["confianza"]), 1) if r["correo"] else None,
            r["confianza_dominio"] if r["dominio"] else None,
            (r["sitio_web"] or "").replace("https://", "").replace("http://", "").lower(),
            " | ".join(r["alternativas"]),
            len(reales),
            " · ".join(f"{ETIQUETA_FORMATO.get(k, k)}: {n}" for k, n in r["desglose"].items()),
            ", ".join(f["correo"] for f in reales if f["formato"])[:1000],
            r["veredicto"],
            " → ".join(ETIQUETA_RUTA.get(x, x) for x in r["ruta"]),
        ])
    ws = wb.create_sheet("Detalle")
    _hoja(ws, ["Cuenta", "Primer nombre", "Segundo nombre", "Primer apellido", "Segundo apellido",
               "Empresa (entrada)", "Empresa encontrada", "Correo estimado", "Patrón", "Dominio",
               "Tipo", "Confianza correo", "Confianza dominio", "Sitio web", "Alternativas",
               "Correos reales en el dominio", "Desglose de formatos", "Correos reales (ejemplos)",
               "Veredicto", "Cómo se construyó"],
          detalle, [14, 16, 16, 16, 16, 30, 30, 40, 24, 26, 40, 10, 10, 26, 60, 10, 50, 60, 90, 90])
    for fila in ws.iter_rows(min_row=2, min_col=19, max_col=20):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(destino)
    return destino


def plantilla(destino: Path | str) -> Path:
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _hoja(wb.active, COLUMNAS, [], [12, 18, 18, 18, 18, 36])
    wb.active.title = "Personas"
    wb.save(destino)
    return destino
